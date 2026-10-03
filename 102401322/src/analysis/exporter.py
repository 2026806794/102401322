"""把统计结果导出成 Excel（作业 2.2 要求用程序自动写入 xlsx）。

用 pandas 组织表格、openpyxl 调样式（表头底色、列宽、冻结首行），
一共导出 13 个工作表。
"""

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.config import XLSX_DIR
from src.utils import get_logger

log = get_logger(__name__)

# 表头样式：深蓝底 + 白字
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
CENTER = Alignment(horizontal="center", vertical="center")
THIN_BORDER = Border(*[Side(style="thin", color="BFBFBF")] * 4)

# 英文键名 -> 中文表头
COLUMN_NAMES = {
    "rank": "排名",
    "name": "名称",
    "word": "词语",
    "count": "数量",
    "ratio": "占比",
    "bvid": "BV号",
    "title": "标题",
    "author": "UP主",
    "keyword": "关键词",
    "view": "播放量",
    "duration": "时长(秒)",
    "danmaku_clean": "有效弹幕数",
    "danmaku_raw": "原始弹幕数",
    "bucket": "视频进度",
}


def _translate_row(row):
    """把一行里的英文键换成中文表头，占比顺便换成百分比。

    占比在数据里是 0.3722 这种小数，写进表格前换成 "37.2%"，
    打开 Excel 看的时候直观一些。
    """
    result = {}
    for key, value in row.items():
        name = COLUMN_NAMES.get(key, key)
        if key == "ratio":
            result[name] = f"{value * 100:.1f}%"
        else:
            result[name] = value
    return result


def _to_dataframe(rows):
    """把字典列表转成 DataFrame。

    转换放在建表之前做（而不是建完再改列），代码更直白，
    也避免了 pandas 的列赋值被静态检查误判成"不支持下标操作"。
    """
    return pd.DataFrame([_translate_row(row) for row in rows])


def _style_sheet(sheet):
    """给一个工作表套样式：表头配色、按内容调列宽、冻结首行。"""
    if sheet.max_row == 0:
        return

    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN_BORDER

    # 列宽按这一列最长内容来定。中文占两个字符宽，所以单独算
    for index, column in enumerate(sheet.iter_cols(min_row=1, max_row=sheet.max_row), start=1):
        max_width = 8
        for cell in column:
            if cell.value is None:
                continue
            width = 0
            for char in str(cell.value):
                if "\u4e00" <= char <= "\u9fff":
                    width += 2
                else:
                    width += 1
            max_width = max(max_width, width)
        width = min(max_width + 4, 60)
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.freeze_panes = "A2"


class ExcelExporter:
    """导出 Excel。"""

    def __init__(self, output_path=None):
        self.output_path = output_path or (XLSX_DIR / "大语言模型弹幕分析统计.xlsx")

    def export(self, result):
        """写文件，返回路径。"""
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        sheets = self._build_sheets(result)

        with pd.ExcelWriter(self.output_path, engine="openpyxl") as writer:
            for name, frame in sheets.items():
                frame.to_excel(writer, sheet_name=name, index=False)
            for name in sheets:
                _style_sheet(writer.sheets[name])

        log.info("Excel 已导出：%s（%d 个工作表）", self.output_path, len(sheets))
        return self.output_path

    def _build_sheets(self, result):
        """组装所有工作表。

        作业要求的三类核心数据（每类弹幕数量、Top8 词频、Top8 案例）放在最前面。
        """
        sheets = {}
        top_n = result.meta.get("top_n", 8)

        # 总览
        overview = [
            {"指标": "学生学号", "数值": result.meta.get("student_id", "")},
            {"指标": "视频总数", "数值": result.total_videos},
            {"指标": "弹幕总数(清洗后)", "数值": result.total_danmaku},
            {"指标": "有效词总数", "数值": result.total_words},
            {"指标": "应用领域类别数", "数值": len(result.category_stats)},
            {"指标": "Top-N 口径", "数值": top_n},
        ]
        for key, value in result.clean_stats.items():
            overview.append({"指标": key, "数值": value})
        sheets["总览"] = _to_dataframe(overview)

        # 1. 每类弹幕的总数量
        sheets["每类弹幕总数量"] = _to_dataframe(result.category_stats)
        # 2. 整体 Top-N 词频
        sheets["Top8弹幕词频"] = _to_dataframe(result.top_words)
        # 3. 排名前 8 的具体应用案例
        sheets["Top8应用案例"] = _to_dataframe(result.case_ranking)

        # 4. 分领域 Top-N 词频（这里只取前 top_n 个，虽然存了 60 个）
        rows = []
        for domain, words in result.top_words_by_category.items():
            for item in words[:top_n]:
                row = {"应用领域": domain}
                row.update(_translate_row(item))
                rows.append(row)
        sheets["分领域Top8词频"] = _to_dataframe(rows)

        sheets["用户态度分布"] = _to_dataframe(result.sentiment)
        sheets["应用成本关注"] = _to_dataframe(result.cost)
        sheets["不利影响关注"] = _to_dataframe(result.risk)

        # 分维度词频（成本/风险/态度各自的高频词）
        rows = []
        for aspect, words in result.top_words_by_aspect.items():
            for item in words[:top_n]:
                row = {"分析维度": aspect}
                row.update(_translate_row(item))
                rows.append(row)
        sheets["分维度Top8词频"] = _to_dataframe(rows)

        sheets["关键词统计"] = _to_dataframe(result.keyword_stats)
        sheets["视频清单"] = _to_dataframe(result.video_stats)
        sheets["弹幕进度分布"] = _to_dataframe(result.progress_distribution)

        rows = []
        for index, item in enumerate(result.word_freq, start=1):
            row = {"rank": index}
            row.update(item)
            rows.append(row)
        sheets["词频明细Top500"] = _to_dataframe(rows)

        return sheets
